#!/usr/bin/env python3
"""tools/master-audit/to_xlsx.py — the flags CSV as a review workbook.

    python3 tools/master-audit/to_xlsx.py <MASTER_AUDIT_<date>_flags.csv> <out.xlsx>

One sheet per kind (check / look) plus "All", alternating gray/white row
bands (Brad's standing rule for anything he reads row by row), a frozen
header, sensible widths, and a "Your decision" column to write in.
Needs openpyxl (pip install openpyxl --break-system-packages).
"""
import csv
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

GRAY = PatternFill('solid', fgColor='EEEEEE')
HEAD = PatternFill('solid', fgColor='1F2A44')
WIDTHS = {'tab': 24, 'item number': 14, 'variation': 9, 'rule': 20, 'kind': 7, 'field': 12,
          'cell as it stands': 40, 'note': 60, 'csv row': 8, 'new': 5, 'Your decision': 28}


def sheet(wb, title, header, rows):
    ws = wb.create_sheet(title)
    ws.append(header + ['Your decision'])
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = HEAD
        c.alignment = Alignment(vertical='center', wrap_text=True)
    for i, r in enumerate(rows):
        ws.append(r + [''])
        if i % 2 == 0:
            for c in ws[ws.max_row]:
                c.fill = GRAY
    for j, h in enumerate(header + ['Your decision'], start=1):
        ws.column_dimensions[get_column_letter(j)].width = WIDTHS.get(h, 16)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    return ws


def main(src, dst):
    with open(src, encoding='utf-8', newline='') as f:
        rd = csv.reader(f)
        header = next(rd)
        rows = [r for r in rd]
    kind_i = header.index('kind')
    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, 'check', header, [r for r in rows if r[kind_i] == 'check'])
    sheet(wb, 'look', header, [r for r in rows if r[kind_i] == 'look'])
    sheet(wb, 'All', header, rows)
    wb.save(dst)
    print('wrote', dst, '-', len(rows), 'flags')


if __name__ == '__main__':
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(2)
    main(sys.argv[1], sys.argv[2])
